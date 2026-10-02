"""Pruebas del frontend en Chrome headless (Playwright) — flujo de aprobaciones y perfiles.

Carga frontend/bitacora-obra.html con una base simulada que imita las reglas del artifact:
  lectura nivel Lector:  contratos, unidades, hitos, oficial
  lectura Colaborador+:  todo lo demás (audit/<cuenta>: solo la propia o Editor)
  escritura Editor:      contratos, unidades, hitos, oficial, permisos, config, credenciales
  escritura propia:      audit/<cuenta>
Escenarios: IF, Validador, Gabinete, Administrador y cuenta sin perfil.

Ejecutar:  python test_frontend.py
"""
import datetime as dt
import hashlib
import json
import os
import pathlib
import sys
import urllib.parse

from playwright.sync_api import sync_playwright

AQUI = pathlib.Path(__file__).resolve().parent
CHROME = os.environ.get('CHROME') or (r'C:\Program Files\Google\Chrome\Application\chrome.exe' if os.name == 'nt' else '/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
PAGINA = (AQUI.parent / 'frontend' / 'bitacora-obra.html').read_text(encoding='utf-8')
XSS = '<img src=x onerror="window.__T.xss++">'
HOY = dt.date.today()


def d(dias):
    return (HOY + dt.timedelta(days=dias)).isoformat()


FAKE_RUNTIME = r"""
window.__T = { writes: [], queries: [], listeners: 0, xss: 0 };
(function(){
  var cfg = window.__TCFG, store = new Map(), listeners = new Set(), idc = 0;
  Object.keys(cfg.seed).forEach(function(k){ store.set(k, JSON.parse(JSON.stringify(cfg.seed[k]))); });
  var NIVEL = { view:0, interact:1, admin:2 }[cfg.level];
  var LECTURA_PUBLICA = ['contratos', 'unidades', 'hitos', 'oficial'];
  var SOLO_EDITOR = ['contratos', 'unidades', 'hitos', 'oficial', 'permisos', 'config', 'credenciales'];
  function clone(x){ return x === undefined ? undefined : JSON.parse(JSON.stringify(x)); }
  function puedeLeer(path){
    var s = path.split('/');
    if(s[0] === 'audit') return s[1] === cfg.uid || NIVEL === 2;
    if(s[0] === 'credenciales') return NIVEL === 2;
    if(LECTURA_PUBLICA.indexOf(s[0]) !== -1) return true;
    return NIVEL >= 1;
  }
  function puedeEscribir(path){
    var s = path.split('/');
    if(SOLO_EDITOR.indexOf(s[0]) !== -1) return NIVEL === 2;
    if(s[0] === 'audit') return s[1] === cfg.uid && NIVEL >= 1;
    return NIVEL >= 1;
  }
  function rechazo(){ var e = new Error('invalid_argument'); e.code = 'invalid_argument'; return Promise.reject(e); }
  function snapDoc(path){ var d = store.get(path), ok = d !== undefined && puedeLeer(path); return { id: path.split('/').pop(), exists: ok, data: function(){ return ok ? clone(d) : undefined; }, metadata:{fromCache:false, hasPendingWrites:false} }; }
  function avisar(){ listeners.forEach(function(l){ setTimeout(l.fire, 0); }); }
  function escuchar(fire){
    var l = { fire: function(){ if(!l.muerto) fire(); } };
    listeners.add(l); window.__T.listeners = listeners.size; setTimeout(l.fire, 0);
    return function(){ l.muerto = true; listeners.delete(l); window.__T.listeners = listeners.size; };
  }
  function cumple(d, f){ if(f[1] === '==') return d[f[0]] === f[2]; if(f[1] === 'in') return f[2].indexOf(d[f[0]]) !== -1; return true; }
  function query(col, filtros, orden, lim){
    function correr(){
      var pref = col + '/', n = col.split('/').length + 1, docs = [];
      store.forEach(function(d, p){ if(p.indexOf(pref) === 0 && p.split('/').length === n && puedeLeer(p) && filtros.every(function(f){ return cumple(d, f); })) docs.push(snapDoc(p)); });
      if(orden) docs.sort(function(a, b){ var x = a.data()[orden[0]], y = b.data()[orden[0]]; if(x === undefined) return 1; if(y === undefined) return -1; return (x < y ? -1 : x > y ? 1 : 0) * (orden[1] === 'desc' ? -1 : 1); });
      if(lim) docs = docs.slice(0, lim);
      return { docs: docs, size: docs.length, empty: !docs.length, docChanges: function(){ return []; }, metadata:{fromCache:false, hasPendingWrites:false} };
    }
    return {
      where: function(f, op, v){ return query(col, filtros.concat([[f, op, v]]), orden, lim); },
      orderBy: function(f, dir){ return query(col, filtros, [f, dir || 'asc'], lim); },
      limit: function(n){ return query(col, filtros, orden, n); },
      get: function(){ return Promise.resolve(correr()); },
      onSnapshot: function(next){ window.__T.queries.push({col: col, filtros: filtros}); return escuchar(function(){ next(correr()); }); }
    };
  }
  function docRef(path){
    return {
      id: path.split('/').pop(), path: path,
      get: function(){ return Promise.resolve(snapDoc(path)); },
      set: function(data){ if(!puedeEscribir(path)) return rechazo(); window.__T.writes.push({op:'set', path:path, data:clone(data)}); store.set(path, clone(data)); avisar(); return Promise.resolve(); },
      update: function(data){ if(!puedeEscribir(path) || !store.has(path)) return rechazo(); window.__T.writes.push({op:'update', path:path, data:clone(data)}); store.set(path, Object.assign(clone(store.get(path)), clone(data))); avisar(); return Promise.resolve(); },
      delete: function(){ window.__T.writes.push({op:'delete', path:path}); store.delete(path); avisar(); return Promise.resolve(); },
      onSnapshot: function(next){ window.__T.queries.push({col: path, filtros: []}); return escuchar(function(){ next(snapDoc(path)); }); },
      collection: function(sub){ return colRef(path + '/' + sub); }
    };
  }
  function colRef(col){
    var q = query(col, [], null, null);
    q.path = col;
    q.doc = function(id){ return docRef(col + '/' + (id || ('auto' + (++idc) + 'x'))); };
    q.add = function(data){ var r = q.doc(); return r.set(data).then(function(){ return r; }); };
    return q;
  }
  var dir = cfg.directory || [];
  var userApi = {
    canEdit: async function(){ return NIVEL === 2; },
    can: async function(){ return NIVEL >= 1; },
    me: async function(){ return { id: cfg.uid, name: cfg.name, email: null, isOwner: false, canEdit: NIVEL === 2, avatarUrl: '', color: '' }; },
    id: async function(){ return cfg.uid; },
    profiles: async function(ids){ var o = {}; [].concat(ids).forEach(function(i){ var p = dir.filter(function(x){ return x.id === i; })[0]; o[i] = { id: i, name: p ? p.name : '', avatarUrl: '', color: '', email: null, isMe: i === cfg.uid, guest: false }; }); return o; },
    search: async function(q){ q = (q || '').toLowerCase(); return dir.filter(function(x){ return x.name.toLowerCase().indexOf(q) !== -1; }).slice(0, 8).map(function(x){ return { id: x.id, name: x.name, avatarUrl: '', color: '', email: null, isMe: false, guest: false }; }); }
  };
  window.claude = { use: async function(n){ return n === 'db' ? { doc: docRef, collection: colRef } : (n === 'user' ? userApi : null); } };
})();
"""


def canon(lista):
    return json.dumps(lista, ensure_ascii=False, separators=(',', ':'))


def sha(texto):
    return hashlib.sha256(texto.encode('utf-8')).hexdigest()


DIRECTORIO = [{'id': 'u_if1', 'name': 'Inés Fiscal'}, {'id': 'u_val', 'name': 'Valentina Jefa'},
              {'id': 'u_if2', 'name': 'Ana Pérez'}, {'id': 'u_admin', 'name': 'Admin Central'}]


def base():
    s = {
        'contratos/c1': {'nombre': 'Red Aeroportuaria Austral', 'codDeaf': 'DEAF-001', 'unidadLabel': 'Aeropuerto', 'inspectorNombre': 'Inés Fiscal', 'correoAprobador': 'jefe.depto@example.org'},
        'contratos/c2': {'nombre': 'Teleférico Bicentenario', 'codDeaf': 'DEAF-002', 'unidadLabel': 'Estación', 'inspectorNombre': 'Otro IF'},
        'unidades/u1': {'contratoId': 'c1', 'nombre': 'Aeródromo ' + XSS, 'activo': True},
        'unidades/u2': {'contratoId': 'c2', 'nombre': 'SECRETO-C2-UNIDAD', 'activo': True},
        'hitos/h1': {'contratoId': 'c1', 'unidadId': 'u1', 'nombre': 'Ingeniería definitiva', 'tipo': 'Ingeniería', 'fechaContractual': d(-60)},
        'hitos/h2': {'contratoId': 'c1', 'unidadId': 'u1', 'nombre': 'Inicio de obras ' + XSS, 'tipo': 'Inicio Físico', 'fechaContractual': d(20)},
        'hitos/h3': {'contratoId': 'c1', 'unidadId': 'u1', 'nombre': 'PSP', 'tipo': 'PSP', 'fechaContractual': d(200)},
        'hitos/h9': {'contratoId': 'c2', 'unidadId': 'u2', 'nombre': 'SECRETO-C2-HITO', 'tipo': 'PSP', 'fechaContractual': d(100)},
        'oficial/avance-u1': {'tipo': 'avance', 'contratoId': 'c1', 'unidadId': 'u1', 'valor': 62, 'fechaReporte': d(-10), 'aprobadoTs': 1790000000000},
        'oficial/hito-h1': {'tipo': 'hito', 'contratoId': 'c1', 'hitoId': 'h1', 'fechaVigente': d(-60), 'fechaReal': d(-55), 'aprobadoTs': 1790000000000},
        'oficial/hist-x1': {'tipo': 'historial', 'contratoId': 'c1', 'campo': 'avance', 'ref': 'u1', 'valorAnterior': 50, 'valorNuevo': 62,
                            'fechaReporte': d(-10), 'nota': 'Nota aprobada ' + XSS, 'autorId': 'u_if1', 'revisorId': 'u_val', 'aprobadoTs': 1790000000000},
        'oficial/alerta-pz': {'tipo': 'alerta', 'contratoId': 'c1', 'problematicaId': 'pz', 'titulo': 'Alerta publicada ' + XSS, 'nivelRiesgo': 'Extremo',
                              'escalamiento': ['Directora'], 'estado': 'escalado', 'publicadoTs': 1790000000000, 'isDeleted': False},
        'oficial/alerta-pq': {'tipo': 'alerta', 'contratoId': 'c1', 'problematicaId': 'pq', 'titulo': 'ALERTA-RETIRADA', 'nivelRiesgo': 'Alto',
                              'escalamiento': [], 'estado': 'escalado', 'publicadoTs': 1790000000000, 'isDeleted': True},
        'solicitudes/sx': {'tipo': 'avance', 'contratoId': 'c2', 'unidadId': 'u2', 'valorPropuesto': 99, 'estado': 'PENDIENTE_APROBACION',
                           'nota': 'SECRETO-C2-SOLICITUD', 'autorId': 'u_otro', 'creadoTs': 1790000000000, 'enviadoTs': 1790000000000},
        'problematicas/p2': {'contratoId': 'c2', 'titulo': 'SECRETO-C2-PB', 'descripcion': 'd', 'estado': 'detectado', 'ts': 4},
        'permisos/u_if1': {'rol': 'if', 'contratos': ['c1'], 'correo': 'ines.if@example.org', 'activo': True},
        'permisos/u_val': {'rol': 'validador', 'contratos': ['c1'], 'activo': True},
    }
    return s


resultados = []


def check(nombre, cond, detalle=''):
    resultados.append((nombre, bool(cond)))
    print(('  OK    ' if cond else '  FALLA ') + nombre + ('' if cond else '  -> ' + str(detalle)[:400]))


def abrir(p, cfg):
    cfg.setdefault('directory', DIRECTORIO)
    tmp = AQUI / '_tmp_prueba.html'
    cfg_js = json.dumps(cfg).replace('</', '<\\/')
    tmp.write_text('<meta charset="utf-8"><script>window.__TCFG = ' + cfg_js + ';</script><script>' + FAKE_RUNTIME
                   + '</script>\n' + PAGINA, encoding='utf-8')
    browser = p.chromium.launch(executable_path=CHROME, headless=True)
    page = browser.new_page(viewport={'width': 1280, 'height': 900})
    errores = []
    page.on('pageerror', lambda e: errores.append(str(e)))
    page.goto(tmp.as_uri())
    return browser, page, errores


def escrituras(page):
    return page.evaluate('window.__T.writes')


def esperar(page, prefijo, n):
    page.wait_for_function('([p, n]) => window.__T.writes.slice(n).some(w => w.path.startsWith(p))', arg=[prefijo, n], timeout=15000)
    return escrituras(page)[n:]


def toast(page):
    page.wait_for_timeout(250)
    return page.inner_text('#toast')


def ingresar(page):
    page.wait_for_selector('#gateBtn', state='visible', timeout=15000)
    page.click('#gateBtn')
    page.wait_for_selector('#appShell', state='visible', timeout=15000)
    page.wait_for_timeout(400)


def tab(page, nombre):
    page.click(f'#tabsNav [data-tab={nombre}]')
    page.wait_for_timeout(250)


def escenario_if(p):
    print('\nA. Inspector Fiscal (Colaborador, contrato c1)')
    seed = base()
    seed['problematicas/p1'] = {'contratoId': 'c1', 'unidadId': 'u1', 'titulo': XSS, 'descripcion': '<script>window.__T.xss++</script>',
                                'escalamiento': 'Directora', 'estado': 'detectado', 'nivelRiesgo': '__proto__', 'ts': 3,
                                'bitacora': [{'texto': XSS, 'autor': '<b>z</b>'}], 'checklist': [{'texto': XSS, 'hecho': False}]}
    seed['solicitudes/sr1'] = {'tipo': 'avance', 'contratoId': 'c1', 'unidadId': 'u1', 'valorPropuesto': 80, 'valorVigente': 62, 'fechaReporte': d(-3), 'nota': 'x',
                               'estado': 'RECHAZADO', 'autorId': 'u_if1', 'creadoTs': 1790000000100, 'enviadoTs': 1790000000100, 'revisorId': 'u_val',
                               'revisadoTs': 1790000000500, 'comentarioRevision': 'Falta respaldo fotográfico ' + XSS}
    seed['solicitudes/sa1'] = {'tipo': 'fecha_programada', 'contratoId': 'c1', 'hitoId': 'h3', 'valorPropuesto': d(210), 'valorVigente': d(200), 'nota': 'Ajuste',
                               'estado': 'APROBADO', 'autorId': 'u_if1', 'creadoTs': 1790000000200, 'enviadoTs': 1790000000200, 'revisorId': 'u_val', 'revisadoTs': 1790000000600}
    browser, page, errores = abrir(p, {'uid': 'u_if1', 'name': 'Inés Fiscal', 'level': 'interact', 'seed': seed})
    try:
        page.wait_for_selector('#gatePerfil', state='visible', timeout=15000)
        txt = page.inner_text('#gatePerfilBox')
        check('Ingreso muestra la cuenta, el perfil y los contratos asignados', 'Inés Fiscal' in txt and 'Inspector Fiscal' in txt and 'Red Aeroportuaria Austral' in txt, txt)
        ingresar(page)
        qs = page.evaluate('window.__T.queries')
        ops = [q for q in qs if q['col'] in ('solicitudes', 'problematicas')]
        check('Datos operativos consultados solo de su contrato', len(ops) == 2 and all(q['filtros'] == [['contratoId', '==', 'c1']] for q in ops), ops)
        check('El IF no consulta permisos, auditoría ni configuración', not any(q['col'].startswith(('permisos', 'audit', 'config')) for q in qs), qs)
        check('Selector "Ver como" oculto para el IF', page.is_hidden('#verComoWrap'))
        tab_ids = page.eval_on_selector_all('#tabsNav .tabd', 'els => els.map(e => e.dataset.tab)')
        check('Pestañas del IF', tab_ids == ['contratos', 'avance', 'hitos', 'solicitudes', 'problematicas', 'resueltas'], tab_ids)
        page.click('.ccard[data-cid=c1] .js-radar')
        page.wait_for_timeout(300)
        pasos = page.eval_on_selector_all('.ccard[data-cid=c1] .rstep', 'els => els.map(e => e.className)')
        check('Radar: Ingeniería cumplida, Inicio Físico crítico, PSP futura, PSD sin hito',
              pasos == ['rstep done', 'rstep crit', 'rstep fut', 'rstep none'], pasos)
        check('Radar crítico muestra días restantes', 'Faltan 20 d' in page.inner_text('.ccard[data-cid=c1] .rstep.crit'))
        check('Desglose del cronograma desplegado al hacer clic', page.locator('.ccard[data-cid=c1] .acc table').count() == 1)
        page.click('.ccard[data-cid=c1] .js-hist')
        page.wait_for_timeout(200)
        check('Bitácora histórica con el cambio aprobado', '50% → 62%' in page.inner_text('.ccard[data-cid=c1] .hist'))

        tab(page, 'avance')
        check('Un solo contrato: no hay selector de contrato activo', page.is_hidden('#avCtrWrap'))
        n = len(escrituras(page))
        page.select_option('#cgAvUnidad', 'u1')
        page.fill('#cgAvPct', '150')
        page.click('#formAvance button[type=submit]')
        check('Avance fuera de rango rechazado', 'entre 0 y 100' in toast(page) and len(escrituras(page)) == n)
        check('Historial de avances al costado: valores aprobados con fecha', '62%' in page.inner_text('#avTimeline') and '50% → 62%' not in page.inner_text('#avTimeline') and '+12 pts' in page.inner_text('#avTimeline'), page.inner_text('#avTimeline'))
        page.fill('#cgAvPct', '70')
        page.fill('#cgAvNota', 'Avance de pavimentos')
        page.click('#formAvance button[type=submit]')
        nuevas = esperar(page, 'solicitudes/', n)
        rec = next(w for w in nuevas if w['path'].startswith('audit/'))
        sol = next(w for w in nuevas if w['path'].startswith('solicitudes/'))
        s = sol['data']
        check('Solicitud creada en revisión, con vigente al proponer y autoría de sesión',
              s['estado'] == 'PENDIENTE_APROBACION' and s['valorPropuesto'] == 70 and s['valorVigente'] == 62 and s['autorId'] == 'u_if1' and s['contratoId'] == 'c1', s)
        esperado = sha(canon(['solicitud', s['tipo'], s['contratoId'], s['unidadId'], s['valorPropuesto'], s['fechaReporte'], s['nota'], s['autorId'], s['creadoTs']]))
        check('Recibo antes de la solicitud y hash compatible con Python', nuevas.index(rec) < nuevas.index(sol) and rec['data']['hash'] == esperado)
        page.wait_for_selector('#modalBack', state='visible')
        href = page.get_attribute('#modalBody .js-mailto', 'href')
        cuerpo = urllib.parse.unquote(href)
        check('Al enviar aparece "Avisar por correo al aprobador" con mailto: al aprobador del contrato', href.startswith('mailto:jefe.depto@example.org?subject=') and 'Avisar por correo' in page.inner_text('#modalBody'), href[:80])
        check('El correo redactado trae contrato, sector, dato, valor propuesto y nota',
              'Contrato: Red Aeroportuaria Austral' in cuerpo and 'Sector: Aeródromo' in cuerpo and 'Avance físico' in cuerpo and '70%' in cuerpo and 'Avance de pavimentos' in cuerpo, cuerpo)
        check('El texto del usuario no se interpreta como HTML en la ventana', page.evaluate('window.__T.xss') == 0)
        page.click('#modalBody .js-modal-cerrar')
        check('La ventana se cierra', page.is_hidden('#modalBack'))
        page.fill('#cgAvPct', '75')
        page.click('#formAvance button[type=submit]')
        check('No permite dos propuestas activas para el mismo dato', 'Ya hay una solicitud activa' in toast(page))

        tab(page, 'hitos')
        check('Gantt con los 3 hitos del contrato', page.locator('.js-g-hito').count() == 3 and page.locator('#ganttWrap .g-hoy').count() == 1)
        check('Gantt: Ingeniería cumplida (real), Inicio Físico con desvío o en plazo según oficial',
              'Cumplido' in page.inner_text('.js-g-hito[data-hid=h1]') and page.locator('.js-g-hito[data-hid=h1] .g-mk.real').count() == 1)
        check('Ya no existe el formulario suelto de hitos', page.locator('#formFecha, #cgHito').count() == 0)
        n = len(escrituras(page))
        page.click('.js-g-hito[data-hid=h2]')
        page.wait_for_selector('#formHito')
        ident = page.inner_text('#modalBody .ident')
        check('Popup identifica contrato, sector y hito', 'Red Aeroportuaria Austral' in ident and 'Aeródromo' in ident and 'Inicio de obras' in ident, ident)
        page.fill('#mhFecha', d(40))
        page.click('#formHito button[type=submit]')
        check('Reprogramar exige justificación', 'Justifica' in toast(page) and len(escrituras(page)) == n)
        page.fill('#mhNota', 'Retraso en permisos')
        page.click('#formHito button[type=submit]')
        nuevas = esperar(page, 'solicitudes/', n)
        sh = next(w for w in nuevas if w['path'].startswith('solicitudes/'))['data']
        check('Desde el Gantt se crea la solicitud de fecha programada en revisión', sh['tipo'] == 'fecha_programada' and sh['estado'] == 'PENDIENTE_APROBACION' and sh['hitoId'] == 'h2' and sh['valorPropuesto'] == d(40), sh)
        page.wait_for_selector('#modalBody .js-mailto')
        cuerpo = urllib.parse.unquote(page.get_attribute('#modalBody .js-mailto', 'href'))
        check('Correo de la fecha: contrato, sector, hito, valor propuesto y justificación', 'Hito a modificar: Inicio de obras' in cuerpo and 'Retraso en permisos' in cuerpo and 'Sector: Aeródromo' in cuerpo, cuerpo)
        page.click('#modalBody .js-modal-cerrar')
        page.click('.js-g-hito[data-hid=h2]')
        page.wait_for_selector('#formHito')
        check('Con una propuesta en revisión, el popup lo avisa', 'Ya hay una propuesta en revisión' in page.inner_text('#mhPend'))
        n = len(escrituras(page))
        page.click('#formHito input[value=fecha_real]')
        page.fill('#mhFecha', d(5))
        page.click('#formHito button[type=submit]')
        check('La fecha real no puede ser futura', 'futura' in toast(page) and len(escrituras(page)) == n)
        page.keyboard.press('Escape')
        check('Escape cierra el popup', page.is_hidden('#modalBack'))

        tab(page, 'solicitudes')
        check('Mis solicitudes muestra días en revisión', 'En revisión hace 0 días' in page.inner_text('#solLista'))
        check('Mis solicitudes ofrece avisar por correo al aprobador', page.locator('#solLista a.js-mailto').count() == 2)
        n = len(escrituras(page))
        page.click('#solLista .sol[data-id] >> nth=0 >> .js-sol-cancelar')
        page.click('#solLista .js-sol-cancelar-si')
        nuevas = esperar(page, 'solicitudes/', n)
        check('Cancelar una solicitud = eliminación lógica (sin borrar)', nuevas[-1]['op'] == 'update' and nuevas[-1]['data']['estado'] == 'ELIMINADO'
              and nuevas[-1]['data']['isDeleted'] is True and not any(w['op'] == 'delete' for w in escrituras(page)), nuevas[-1])
        tab(page, 'avance')
        n = len(escrituras(page))
        page.select_option('#cgAvUnidad', 'u1')
        page.fill('#cgAvPct', '66')
        page.click('#formAvance [data-borrador]')
        nuevas = esperar(page, 'solicitudes/', n)
        check('Borrador de avance guardado', next(w for w in nuevas if w['path'].startswith('solicitudes/'))['data']['estado'] == 'BORRADOR')

        tab(page, 'problematicas')
        page.click('.kcard[data-id=p1] .kcard-expand-btn')
        page.wait_for_timeout(200)

        n = len(escrituras(page))
        page.select_option('.kcard[data-id=p1] .k-move', 'resuelto')
        page.wait_for_selector('#formResolver')
        page.click('#formResolver button[type=submit]')
        check('Resolver exige el motivo (sin motivo no se guarda)', 'mínimo 5' in toast(page) and len(escrituras(page)) == n)
        page.fill('#mrTexto', 'Se acordó el plan de acceso con la comunidad')
        page.click('#formResolver button[type=submit]')
        nuevas = esperar(page, 'problematicas/p1', n)
        up = nuevas[-1]['data']
        check('Resolución guardada con texto, fecha y autor (id, no nombre)', up['estado'] == 'resuelto' and up['resolucion']['texto'].startswith('Se acordó') and up['resolucion']['autorId'] == 'u_if1'
              and up['resolucion']['fecha'] == up['fechaResuelta'] and 'Inés' not in json.dumps(up), up)
        tab(page, 'resueltas')
        check('Resueltas muestra el motivo', 'Se acordó el plan' in page.inner_text('#resueltasBody'))
        page.click('#resueltasBody tr[data-id=p1] .btn-ficha')
        page.wait_for_selector('#modalBody .ficha-sec')
        ficha = page.inner_text('#modalBody').lower()
        check('Ficha completa: descripción, checklist, bitácora y motivo con fecha y autor',
              'qué es el problema' in ficha and 'tareas para subsanar' in ficha and 'bitácora de notas (1)' in ficha
              and 'se acordó el plan de acceso' in ficha and 'inés fiscal (tú)' in ficha and 'resuelta el' in ficha, ficha)
        page.click('#modalBody .js-modal-cerrar')
        n = len(escrituras(page))
        page.click('#resueltasBody tr[data-id=p1] .btn-reabrir')
        nuevas = esperar(page, 'problematicas/p1', n)
        up = nuevas[-1]['data']
        check('Reabrir conserva la resolución anterior (nada se pierde)', up['estado'] == 'detectado' and up['resolucion'] is None and len(up['resolucionesPrevias']) == 1, up)

        check('Campanita: 2 avisos sin leer (1 rechazada y 1 aprobada)', page.inner_text('#bellN') == '2', page.inner_text('#bellN'))
        page.click('#btnNotif')
        panel = page.inner_text('#notifPanel').lower()
        check('Panel de notificaciones: rechazadas con comentario, aprobadas y pendientes', 'rechazadas (1)' in panel and 'falta respaldo fotográfico' in panel and 'aprobadas (1)' in panel and 'pendientes de revisión (1)' in panel, panel)
        n = len(escrituras(page))
        page.click('#notifLeidas')
        nuevas = esperar(page, 'notificaciones/u_if1', n)
        page.wait_for_selector('#bellN', state='hidden', timeout=5000)
        check('Marcar leídas guarda leidoTs por usuario y apaga el contador', isinstance(nuevas[-1]['data']['leidoTs'], int) and set(nuevas[-1]['data']) == {'leidoTs'} and page.is_hidden('#bellN'), nuevas[-1])
        page.click('#notifIr')
        check('El panel lleva a Mis solicitudes (resueltas)', page.locator('#solFiltros .filter-chip.active', has_text='Resueltas').count() == 1 and page.is_hidden('#notifPanel'))
        check('Ninguna carga XSS se ejecutó en ninguna vista', page.evaluate('window.__T.xss') == 0, page.evaluate('window.__T.xss'))
        check('Ningún dato de otro contrato llega a la página', 'SECRETO' not in page.evaluate('document.body.innerHTML'))
        check('El IF no ve insignias de verificación', page.locator('.vbadge').count() == 0)
        page.set_viewport_size({'width': 1920, 'height': 900})
        page.wait_for_timeout(200)
        ancho = page.evaluate("document.querySelector('.content-wrap').getBoundingClientRect().width")
        logo = page.evaluate("document.querySelector('.appbar-logo').getBoundingClientRect().height")
        check('Pantalla ancha: el contenido crece más allá de 1.180 px (máx. 1.600)', 1180 < ancho <= 1600, ancho)
        check('Logo del banner grande y legible (>= 60 px)', logo >= 60, logo)
        check('Banner y contenido usan el mismo ancho', page.evaluate("document.querySelector('.appbar-inner').getBoundingClientRect().width") == ancho)
        page.set_viewport_size({'width': 390, 'height': 800})
        page.wait_for_timeout(200)
        check('Ancho de teléfono: sin scroll horizontal de la página', page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), page.evaluate('document.documentElement.scrollWidth'))
        page.set_viewport_size({'width': 1280, 'height': 900})
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_if_dos_contratos(p):
    print('\nB0. IF con dos contratos (titular y suplente)')
    seed = base()
    for k in [k for k in seed if 'SECRETO' in json.dumps(seed[k])]:
        pass
    seed['contratos/c3'] = {'nombre': 'Hospital Norte', 'unidadLabel': 'Hospital', 'inspectorNombre': 'IF titular'}
    seed['unidades/u3'] = {'contratoId': 'c3', 'nombre': 'Torre Hospital', 'activo': True}
    seed['hitos/h4'] = {'contratoId': 'c3', 'unidadId': 'u3', 'nombre': 'Hito Hospital 1', 'tipo': 'Ingeniería', 'fechaContractual': d(30)}
    seed['hitos/h5'] = {'contratoId': 'c3', 'unidadId': 'u3', 'nombre': 'Hito Hospital 2', 'tipo': 'PSP', 'fechaContractual': d(300)}
    seed['permisos/u_if1'] = {'rol': 'if', 'contratos': ['c1', 'c3'], 'rolesContrato': {'c1': 'titular', 'c3': 'suplente'}, 'activo': True}
    browser, page, errores = abrir(p, {'uid': 'u_if1', 'name': 'Inés Fiscal', 'level': 'interact', 'seed': seed})
    try:
        ingresar(page)
        check('Tarjetas con etiqueta Titular / Suplente', 'Titular' in page.inner_text('.ccard[data-cid=c1] h3') and 'Suplente' in page.inner_text('.ccard[data-cid=c3] h3'),
              [page.inner_text('.ccard[data-cid=c1] h3'), page.inner_text('.ccard[data-cid=c3] h3')])
        tab(page, 'avance')
        opts = page.eval_on_selector_all('#avContrato option', 'els => els.map(e => e.textContent)')
        check('Con más de un contrato aparece el selector de contrato activo (con Titular/Suplente)', page.is_visible('#avCtrWrap') and len(opts) == 2 and any('Titular' in o for o in opts) and any('Suplente' in o for o in opts), opts)
        check('Unidades del contrato activo solamente', page.eval_on_selector_all('#cgAvUnidad option', 'els => els.map(e => e.value)') == ['u3'])
        page.select_option('#avContrato', 'c1')
        page.wait_for_timeout(250)
        check('Al cambiar de contrato cambian las unidades', page.eval_on_selector_all('#cgAvUnidad option', 'els => els.map(e => e.value)') == ['u1'])
        page.select_option('#avContrato', 'c3')
        tab(page, 'hitos')
        check('El Gantt sigue el contrato activo (Hospital Norte: 2 hitos)', page.locator('.js-g-hito').count() == 2 and 'Hito Hospital 1' in page.inner_text('#ganttWrap'))
        page.select_option('#htContrato', 'c1')
        page.wait_for_timeout(250)
        check('Volver al otro contrato muestra sus 3 hitos', page.locator('.js-g-hito').count() == 3)
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_validador(p):
    print('\nB. Jefe de Departamento / Validador (Editor, contrato c1)')
    seed = base()
    s1 = {'tipo': 'avance', 'contratoId': 'c1', 'unidadId': 'u1', 'valorPropuesto': 70, 'valorVigente': 62, 'fechaReporte': d(-1),
          'nota': 'Pavimentos ' + XSS, 'estado': 'PENDIENTE_APROBACION', 'autorId': 'u_if1', 'creadoTs': 1790000000001, 'enviadoTs': 1790000000001}
    seed['solicitudes/s1'] = s1
    seed['audit/u_if1/eventos/solicitud-s1-crear'] = {'hash': sha(canon(['solicitud', 'avance', 'c1', 'u1', 70, s1['fechaReporte'], s1['nota'], 'u_if1', 1790000000001]))}
    seed['solicitudes/s2'] = {'tipo': 'fecha_programada', 'contratoId': 'c1', 'hitoId': 'h2', 'valorPropuesto': d(45), 'nota': 'Permisos',
                              'estado': 'PENDIENTE_APROBACION', 'autorId': 'u_if1', 'creadoTs': 1790000000002, 'enviadoTs': 1790000000002}
    seed['solicitudes/s3'] = {'tipo': 'fecha_real', 'contratoId': 'c1', 'hitoId': 'h3', 'valorPropuesto': d(-1), 'estado': 'PENDIENTE_APROBACION',
                              'autorId': 'u_val', 'creadoTs': 1790000000003, 'enviadoTs': 1790000000003}
    seed['problematicas/p1'] = {'contratoId': 'c1', 'unidadId': 'u1', 'titulo': 'Toma de terreno', 'descripcion': 'x', 'estado': 'escalado',
                                'nivelRiesgo': 'Extremo', 'escalamiento': ['Directora'], 'ts': 5, 'autorId': 'u_if1'}
    browser, page, errores = abrir(p, {'uid': 'u_val', 'name': 'Valentina Jefa', 'level': 'admin', 'seed': seed})
    try:
        ingresar(page)
        check('Badge con pendientes en la pestaña Bandeja', page.inner_text('#tabsNav [data-tab=bandeja] .tab-badge') == '3')
        tab(page, 'bandeja')
        page.wait_for_function("document.querySelector('#bjLista .sol[data-id=s1] .vbadge') && !document.querySelector('#bjLista .sol[data-id=s1] .vbadge').textContent.includes('verificando')", timeout=10000)
        card = page.inner_text('#bjLista .sol[data-id=s1]')
        check('Comparativo valor actual vs propuesto con diferencia', '62%' in card and '70%' in card and '+8 puntos porcentuales' in card, card)
        check('Autoría de la solicitud verificada por recibo', '✓ verificado' in card, card)
        ident = page.inner_text('#bjLista .sol[data-id=s2] .ident').lower()
        check('Solicitud de fecha: Contrato / Sector / Hito a modificar bien identificados',
              'contrato' in ident and 'red aeroportuaria austral' in ident and 'sector' in ident and 'aeródromo' in ident and 'hito a modificar' in ident and 'inicio de obras' in ident, ident)
        check('Esos datos se ven con letra grande', page.evaluate("parseFloat(getComputedStyle(document.querySelector('#bjLista .sol[data-id=s2] .ident strong')).fontSize)") >= 17)
        s2txt = page.inner_text('#bjLista .sol[data-id=s2]')
        check('Muestra valor actual → propuesto y la justificación del IF', 'VALOR ACTUAL' in s2txt.upper() and 'VALOR PROPUESTO' in s2txt.upper(), s2txt)
        check('Justificación visible', 'Permisos' in s2txt and 'Justificación del IF' in s2txt, s2txt)
        check('Solicitud propia del validador no se puede aprobar', page.locator('#bjLista .sol[data-id=s3] .js-aprobar').count() == 0)
        check('Solicitudes de otros contratos no aparecen', page.locator('#bjLista .sol[data-id=sx]').count() == 0)

        n = len(escrituras(page))
        page.click('#bjLista .sol[data-id=s1] .js-aprobar')
        nuevas = esperar(page, 'solicitudes/s1', n)
        rutas = [w['path'] for w in nuevas]
        of = next(w for w in nuevas if w['path'] == 'oficial/avance-u1')['data']
        check('Aprobar publica el valor oficial, el historial y cierra la solicitud',
              of['valor'] == 70 and 'oficial/hist-s1' in rutas and nuevas[-1]['data']['estado'] == 'APROBADO', rutas)
        page.wait_for_selector('#modalBody .js-mailto')
        cuerpo = urllib.parse.unquote(page.get_attribute('#modalBody .js-mailto', 'href'))
        check('Tras aprobar, "Avisar por correo al IF" va dirigido al IF autor', cuerpo.startswith('mailto:ines.if@example.org?subject=') and 'APROBADA' in cuerpo and 'Contrato: Red Aeroportuaria Austral' in cuerpo and '70%' in cuerpo, cuerpo)
        page.click('#modalBody .js-modal-cerrar')

        n = len(escrituras(page))
        page.click('#bjLista .sol[data-id=s2] .js-rechazar')
        page.click('#bjLista .sol[data-id=s2] .js-rechazo-ok')
        check('Rechazar sin comentario no se permite', 'obligatorio' in toast(page) and len(escrituras(page)) == n)
        page.fill('#bjLista .sol[data-id=s2] .bj-motivo', 'Falta respaldo del permiso sectorial')
        page.click('#bjLista .sol[data-id=s2] .js-rechazo-ok')
        nuevas = esperar(page, 'solicitudes/s2', n)
        check('Rechazo guardado con comentario', nuevas[-1]['data']['estado'] == 'RECHAZADO' and 'respaldo' in nuevas[-1]['data']['comentarioRevision'])
        page.wait_for_selector('#modalBody .js-mailto')
        cuerpo = urllib.parse.unquote(page.get_attribute('#modalBody .js-mailto', 'href'))
        check('Tras rechazar, el correo al IF incluye el comentario del revisor', cuerpo.startswith('mailto:ines.if@example.org') and 'RECHAZADA' in cuerpo and 'Motivo del rechazo: Falta respaldo del permiso sectorial' in cuerpo and 'Sector: Aeródromo' in cuerpo and 'Hito a modificar: Inicio de obras' in cuerpo, cuerpo)
        page.click('#modalBody .js-modal-cerrar')

        n = len(escrituras(page))
        page.click('#bjAlertas .js-publicar[data-pid=p1]')
        nuevas = esperar(page, 'oficial/alerta-p1', n)
        check('Publicar alerta para Gabinete', nuevas[-1]['data']['isDeleted'] is False and nuevas[-1]['data']['nivelRiesgo'] == 'Extremo')
        n = len(escrituras(page))
        page.wait_for_selector('#bjAlertas .js-retirar[data-pid=p1]')
        page.click('#bjAlertas .js-retirar[data-pid=p1]')
        nuevas = esperar(page, 'oficial/alerta-p1', n)
        check('Retirar alerta = eliminación lógica', nuevas[-1]['op'] == 'set' and nuevas[-1]['data']['isDeleted'] is True)

        check('Selector "Ver como" visible para Editores', page.is_visible('#verComoWrap'))
        nq = len(page.evaluate('window.__T.queries'))
        page.select_option('#verComo', 'gabinete')
        page.wait_for_timeout(500)
        cols = [q['col'] for q in page.evaluate('window.__T.queries')[nq:]]
        check('Vista previa Gabinete solo consulta colecciones oficiales', sorted(cols) == ['hitos', 'oficial', 'unidades'], cols)
        check('Aviso de vista previa visible', 'Vista previa como Gabinete' in page.inner_text('#bannerTxt'))
        check('Ninguna carga XSS se ejecutó', page.evaluate('window.__T.xss') == 0)
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_gabinete(p):
    print('\nC. Gabinete (Lector: solo datos aprobados)')
    seed = base()
    seed['solicitudes/sp'] = {'tipo': 'avance', 'contratoId': 'c1', 'unidadId': 'u1', 'valorPropuesto': 99, 'estado': 'PENDIENTE_APROBACION',
                              'nota': 'NO-APROBADO', 'autorId': 'u_if1', 'creadoTs': 1790000000009, 'enviadoTs': 1790000000009}
    seed['oficial/hist-h2'] = {'tipo': 'historial', 'contratoId': 'c1', 'campo': 'fecha_programada', 'ref': 'h2', 'valorAnterior': d(20), 'valorNuevo': d(35),
                               'fechaReporte': d(-1), 'nota': 'Permisos', 'autorId': 'u_if1', 'revisorId': 'u_val', 'aprobadoTs': 1790500000000}
    for i in range(11):
        seed[f'oficial/hist-a{i:02d}'] = {'tipo': 'historial', 'contratoId': 'c1', 'campo': 'avance', 'ref': 'u1', 'valorAnterior': i, 'valorNuevo': i + 1,
                                          'fechaReporte': d(-30), 'nota': '', 'autorId': 'u_if1', 'revisorId': 'u_val', 'aprobadoTs': 1780000000000 + i * 1000}
    browser, page, errores = abrir(p, {'uid': 'u_gab', 'name': 'Gabinete DG', 'level': 'view', 'seed': seed})
    try:
        page.wait_for_selector('#gatePerfil', state='visible', timeout=15000)
        check('Lector entra automáticamente como Gabinete', 'Gabinete' in page.inner_text('#gatePerfilBox'))
        ingresar(page)
        cols = sorted(q['col'] for q in page.evaluate('window.__T.queries'))
        check('Gabinete solo consulta colecciones oficiales', cols == ['contratos', 'hitos', 'oficial', 'unidades'], cols)
        cuerpo = page.inner_text('#panel-gabinete')
        check('Muestra el avance aprobado (62%) y no la propuesta pendiente (99%)', '62%' in cuerpo and '99%' not in cuerpo and 'NO-APROBADO' not in cuerpo)
        check('Lista la alerta publicada y no la retirada', 'Alerta publicada' in cuerpo and 'ALERTA-RETIRADA' not in cuerpo)
        check('Semáforo rojo por alerta extrema', page.inner_text('.ccard[data-cid=c1] .sem') == 'Rojo')
        check('Sin pesos definidos usa promedio simple y lo dice', 'promedio simple' in page.inner_text('.ccard[data-cid=c1]'))
        n_rojo, n_amb, n_verde = [page.inner_text(f'#gbKpis .sem-fig.{c} .sem-n') for c in ('bad', 'warn', 'good')]
        check('Semáforo: tres cifras con color y etiqueta ("1 contrato en rojo")', (n_rojo, n_amb, n_verde) == ('1', '0', '1') and '1 contrato en rojo' in ' '.join(page.inner_text('#gbKpis .sem-fig.bad').split())
              and '0 contratos en ámbar' in ' '.join(page.inner_text('#gbKpis .sem-fig.warn').split()) and '1 contrato en verde' in ' '.join(page.inner_text('#gbKpis .sem-fig.good').split()), (n_rojo, n_amb, n_verde))
        ley = page.inner_text('#gbKpis .sem-ley')
        check('Leyenda de cómo se calcula (rojo, ámbar y verde)', 'Rojo' in ley and 'hito atrasado' in ley and 'alerta extrema' in ley and 'Ámbar' in ley and 'desvío' in ley and 'Verde' in ley and 'ninguna' in ley, ley)
        check('El "2/0/1" confuso ya no existe', '1/0/1' not in page.inner_text('#gbKpis') and page.locator('#gbKpis .kpi-num', has_text='/').count() == 0)
        page.click('#gbKpis .sem-fig.bad')
        check('Clic en rojo: solo el contrato en rojo, y la cifra no cambia', page.locator('#gbCards .ccard').count() == 1 and page.locator('#gbCards .ccard[data-cid=c1]').count() == 1
              and page.get_attribute('#gbKpis .sem-fig.bad', 'aria-pressed') == 'true' and page.inner_text('#gbKpis .sem-fig.good .sem-n') == '1')
        page.click('#gbKpis .sem-fig.good')
        check('Clic en verde: solo el contrato en verde', page.locator('#gbCards .ccard').count() == 1 and page.locator('#gbCards .ccard[data-cid=c2]').count() == 1)
        page.click('#gbKpis .sem-fig.warn')
        check('Un color sin contratos lo dice', page.locator('#gbCards .ccard').count() == 0 and 'Ningún contrato en ámbar' in page.inner_text('#gbEmpty'))
        page.click('#gbKpis .js-sem-quitar')
        check('Quitar el filtro muestra todos', page.locator('#gbCards .ccard').count() == 2)
        page.click('#gbKpis .sem-fig.bad')
        page.click('#gbKpis .sem-fig.bad')
        check('Volver a hacer clic en el mismo color quita el filtro', page.locator('#gbCards .ccard').count() == 2)
        check('Sin formularios ni botones de edición', page.locator('#panel-gabinete input, #panel-gabinete textarea, #panel-gabinete .js-aprobar').count() == 0)
        check('Gabinete no tiene bandeja ni botones de aprobación', page.locator('.js-aprobar, .js-rechazar, .js-publicar, .js-retirar, #tabsNav [data-tab=bandeja], #btnNotif:visible').count() == 0 and page.is_hidden('#panel-bandeja'))
        filas = page.locator('#gbAprobaciones tbody tr')
        check('Últimas aprobaciones: muestra 10 y ofrece "Ver más"', filas.count() == 10 and page.is_visible('#gbVerMas'), filas.count())
        primera = filas.nth(0).inner_text()
        fecha_apr = page.evaluate("new Date(1790500000000).toLocaleDateString('es-CL', {day:'2-digit', month:'short', year:'numeric'})")
        check('La más reciente es el hito modificado: contrato, sector, hito, fecha anterior → nueva y fecha de aprobación',
              'Red Aeroportuaria Austral' in primera and 'Aeródromo' in primera and 'Inicio de obras' in primera and 'Fecha programada' in primera and '→' in primera and fecha_apr in primera, primera)
        check('Orden de la más reciente a la más antigua', 'Fecha programada' in filas.nth(0).inner_text() and 'Avance físico' in filas.nth(1).inner_text() and '62%' in filas.nth(1).inner_text(), filas.nth(1).inner_text())
        page.click('#gbVerMas')
        check('"Ver más" muestra el resto y se oculta', page.locator('#gbAprobaciones tbody tr').count() == 13 and page.is_hidden('#gbVerMas'))
        page.click('#gbFiltros [data-filtro=hito]')
        check('Filtro "Hitos modificados"', page.locator('#gbAprobaciones tbody tr').count() == 1)
        page.click('#gbFiltros [data-filtro=avance]')
        check('Filtro "Avances"', page.locator('#gbAprobaciones tbody tr').count() == 10 and 'Fecha programada' not in page.inner_text('#gbAprobaciones'))
        page.click('#gbFiltros [data-filtro=todo]')
        check('Una sola pestaña y sin selector de perfil', page.locator('#tabsNav .tabd').count() == 1 and page.is_hidden('#verComoWrap'))
        check('Ninguna carga XSS se ejecutó', page.evaluate('window.__T.xss') == 0)
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_admin(p):
    print('\nD. Administrador del sistema (Editor sin perfil explícito)')
    browser, page, errores = abrir(p, {'uid': 'u_admin', 'name': 'Admin Central', 'level': 'admin', 'seed': base()})
    try:
        ingresar(page)
        tab(page, 'permisos')
        page.click('#permBuscar')
        page.fill('#permBuscar', 'Ana')
        page.wait_for_selector('#permResultados .js-elegir[data-uid=u_if2]')
        page.click('#permResultados .js-elegir[data-uid=u_if2]')
        page.wait_for_selector('#permTabla tr[data-uid=u_if2]')
        n = len(escrituras(page))
        page.click('#permTabla tr[data-uid=u_if2] .js-guardar-perm')
        check('Perfil IF sin contratos no se guarda', 'al menos un contrato' in toast(page) and len(escrituras(page)) == n)
        page.click('#permTabla tr[data-uid=u_if2] .perm-chip[data-cid=c2]')
        page.select_option('#permTabla tr[data-uid=u_if2] .f-rolctr[data-cid=c2]', 'suplente')
        page.fill('#permTabla tr[data-uid=u_if2] .f-correo', 'correo-malo')
        page.click('#permTabla tr[data-uid=u_if2] .js-guardar-perm')
        check('Un correo con formato inválido no se guarda', 'formato' in toast(page) and len(escrituras(page)) == n)
        page.fill('#permTabla tr[data-uid=u_if2] .f-correo', 'ana.if@example.org')
        page.click('#permTabla tr[data-uid=u_if2] .js-guardar-perm')
        nuevas = esperar(page, 'permisos/u_if2', n)
        dd = nuevas[-1]['data']
        check('Permiso guardado con multiselección de contratos', dd['rol'] == 'if' and dd['contratos'] == ['c2'] and dd['activo'] is True, dd)
        check('Permiso guarda Titular/Suplente por contrato y el correo de la persona', dd['rolesContrato'] == {'c2': 'suplente'} and dd['correo'] == 'ana.if@example.org', dd)
        n = len(escrituras(page))
        page.fill('#permAprobadores tr[data-cid=c1] .f-correo-apr', 'aprobador.c1@example.org')
        page.click('#permAprobadores tr[data-cid=c1] .js-guardar-apr')
        nuevas = esperar(page, 'contratos/c1', n)
        check('Correo del aprobador guardado en el contrato', nuevas[-1]['op'] == 'update' and nuevas[-1]['data'] == {'correoAprobador': 'aprobador.c1@example.org'}, nuevas[-1])
        check('Solo se guarda el id de la persona, no su nombre', 'Ana' not in json.dumps(dd))

        tab(page, 'config')
        n = len(escrituras(page))
        page.click('#cfHitos tr[data-id=h3] .js-hito-del')
        page.click('#cfHitos tr[data-id=h3] .js-hito-del-si')
        nuevas = esperar(page, 'hitos/h3', n)
        check('Eliminar hito = marca isDeleted (no borra)', nuevas[-1]['op'] == 'set' and nuevas[-1]['data']['isDeleted'] is True)
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_cargas_admin(p):
    print('\nD2. Administrador: cargas por Excel/CSV y subsectores por contrato')
    import openpyxl
    tmpd = AQUI / '_tmp_cargas'
    tmpd.mkdir(exist_ok=True)
    CAB = ['COD_DEAF', 'NOMBRE_CONTRATO', 'TIPO_UNIDAD', 'IF_TITULAR', 'IF_SUPLENTE', 'CORREO_APROBADOR']

    def xlsx(nombre, filas, cab=CAB):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Contratos'
        ws.append(cab)
        for f in filas:
            ws.append(f)
        ruta = tmpd / nombre
        wb.save(ruta)
        return ruta

    browser, page, errores = abrir(p, {'uid': 'u_admin', 'name': 'Admin Central', 'level': 'admin', 'seed': base()})
    try:
        ingresar(page)
        tab(page, 'config')
        # --- plantilla generada por la página, leída con una librería externa ---
        with page.expect_download() as dl:
            page.click('.js-plantilla[data-tipo=contratos][data-formato=xlsx]')
        ruta = tmpd / 'plantilla.xlsx'
        dl.value.save_as(ruta)
        wb = openpyxl.load_workbook(ruta)
        check('Plantilla .xlsx de contratos: abre en un lector externo con las 6 columnas', wb.sheetnames == ['Contratos', 'Instrucciones'] and [c.value for c in wb['Contratos'][1]] == CAB, wb.sheetnames)
        with page.expect_download() as dl:
            page.click('.js-plantilla[data-tipo=contratos][data-formato=csv]')
        csv_txt = dl.value.path().read_bytes().decode('utf-8-sig')
        check('Plantilla .csv con las mismas columnas', csv_txt.strip().split(';') == CAB, csv_txt)
        page.set_input_files('#impContratosFile', str(ruta))
        page.wait_for_selector('#impContratosBox .sol-rechazo')
        check('Subir la plantilla vacía avisa que no hay filas', 'no tiene filas con datos' in page.inner_text('#impContratosBox'), page.inner_text('#impContratosBox'))

        # --- archivo con errores: nada se importa y se ve cada error por fila ---
        n = len(escrituras(page))
        malo = xlsx('malo.xlsx', [
            ['DEAF-003', 'Nuevo Contrato ' + XSS, 'Sector', 'IF Titular', '', 'nuevo@example.org'],
            ['DEAF-001', 'Red Aeroportuaria Austral II', 'Aeropuerto', 'IF', '', 'x@example.org'],
            ['DEAF-003', 'Duplicado', 'Sector', 'IF', '', 'd@example.org'],
            ['', 'Sin código', 'Sector', 'IF', '', 'e@example.org'],
            ['DEAF-005', 'Correo malo', 'Sector', 'IF', '', 'no-es-correo'],
            ['DEAF-006', "=cmd|' /C calc'!A0", 'Sector', 'IF', '', 'f@example.org'],
            ['DEAF-007', 'Sin correo', 'Sector', 'IF', '', ''],
            ['DEAF 008!', 'Código raro', 'Sector', 'IF', '', 'g@example.org'],
            ['DEAF-009', '+cmd', 'Sector', 'IF', '', 'h@example.org'],
        ])
        page.set_input_files('#impContratosFile', str(malo))
        page.wait_for_selector('#impContratosBox table')
        box = page.inner_text('#impContratosBox')
        check('Vista previa: 7 filas con error contadas', '7 filas con errores' in box, box[:300])
        for texto in ('COD_DEAF repetido en el archivo', 'COD_DEAF es obligatorio', 'CORREO_APROBADOR no es un correo válido', 'no puede comenzar con = + - @',
                      'CORREO_APROBADOR es obligatorio', 'COD_DEAF solo admite letras'):
            check('Error por fila: ' + texto, texto in box, box)
        check('Errores indicados por número de fila de la hoja', page.locator('#impContratosBox tbody tr').nth(2).inner_text().startswith('4'), page.locator('#impContratosBox tbody tr').nth(2).inner_text())
        check('Con errores, el botón de confirmar está bloqueado y no se escribió nada', page.is_disabled('#impContratosBox .js-imp-ok') and len(escrituras(page)) == n)
        check('Contenido de las celdas escapado en la vista previa (sin XSS)', page.evaluate('window.__T.xss') == 0)
        page.click('#impContratosBox .js-imp-no')
        check('Cancelar limpia la vista previa', page.inner_text('#impContratosBox').strip() == '')

        # --- archivo correcto: crea y actualiza por COD_DEAF; nada se borra ---
        bueno = xlsx('bueno.xlsx', [
            ['DEAF-003', 'Nuevo Contrato', 'Sector', 'IF Titular', 'IF Suplente', 'nuevo@example.org'],
            ['deaf-001', 'Red Aeroportuaria Austral II', 'Aeropuerto', 'IF', '', 'x@example.org'],
            [8, 'Contrato con código numérico', 'Hospital', 'IF', '', 'h@example.org'],
            ['DEAF-002', 'Teleférico Bicentenario', 'Estación', 'Otro IF', '', ''],
        ])
        page.set_input_files('#impContratosFile', str(bueno))
        page.wait_for_selector('#impContratosBox table')
        box = page.inner_text('#impContratosBox')
        check('Hay errores solo en la fila sin correo (DEAF-002 sigue sin correo aprobador)', '1 fila con errores' in box, box[:200])
        bueno = xlsx('bueno.xlsx', [
            ['DEAF-003', 'Nuevo Contrato', 'Sector', 'IF Titular', 'IF Suplente', 'nuevo@example.org'],
            ['deaf-001', 'Red Aeroportuaria Austral II', 'Aeropuerto', 'IF', '', 'x@example.org'],
            [8, 'Contrato con código numérico', 'Hospital', 'IF', '', 'h@example.org'],
        ])
        page.set_input_files('#impContratosFile', str(bueno))
        page.wait_for_selector('#impContratosBox table')
        box = page.inner_text('#impContratosBox')
        check('Vista previa: 2 nuevos, 1 por actualizar (el COD_DEAF se compara sin distinguir mayúsculas)', '2 nuevos' in box and '1 por actualizar' in box and 'errores' not in box, box[:300])
        n = len(escrituras(page))
        page.click('#impContratosBox .js-imp-ok')
        page.wait_for_function('([n]) => window.__T.writes.length >= n + 3', arg=[n], timeout=10000)
        nuevas = escrituras(page)[n:]
        upd = [w for w in nuevas if w['op'] == 'update']
        crea = [w for w in nuevas if w['op'] == 'set' and w['path'].startswith('contratos/')]
        check('Se actualiza el contrato existente por COD_DEAF sin tocar su estado "activo"', len(upd) == 1 and upd[0]['path'] == 'contratos/c1' and upd[0]['data']['nombre'] == 'Red Aeroportuaria Austral II' and 'activo' not in upd[0]['data'], upd)
        check('Se crean los contratos nuevos con codDeaf (numérico como texto)', sorted(w['data']['codDeaf'] for w in crea) == ['8', 'DEAF-003'] and all(w['data']['activo'] is True for w in crea), crea)
        check('Nada se borra y el otro contrato queda intacto', not any(w['op'] == 'delete' or w['path'] == 'contratos/c2' for w in nuevas))
        check('Aviso de importación lista', 'Importación lista' in toast(page))
        page.wait_for_selector('#cfContratos tr input.f-coddeaf[value="DEAF-003"]')
        check('Los contratos importados aparecen en la tabla con su COD_DEAF', page.locator('#cfContratos tr').count() == 4)

        # --- CSV de Excel en español: separador ';', comillas y BOM ---
        csv = tmpd / 'contratos.csv'
        csv.write_bytes(('\ufeffNOMBRE_CONTRATO;COD_DEAF;correo aprobador;IF_TITULAR;Tipo unidad\r\n"Contrato ""Norte""; con punto y coma";DEAF-020;n@example.org;IF Norte;Sector\r\n').encode('utf-8'))
        page.set_input_files('#impContratosFile', str(csv))
        page.wait_for_selector('#impContratosBox table')
        box = page.inner_text('#impContratosBox')
        check('CSV: reconoce separador ";", comillas, BOM y columnas en otro orden/nombre', 'Contrato "Norte"; con punto y coma' in box and '1 nuevo' in box and 'errores' not in box, box[:300])
        page.click('#impContratosBox .js-imp-no')

        # --- archivos inválidos ---
        txt = tmpd / 'x.txt'; txt.write_text('hola')
        page.set_input_files('#impContratosFile', str(txt))
        page.wait_for_selector('#impContratosBox .sol-rechazo')
        check('Formato no permitido', 'Formato no permitido' in page.inner_text('#impContratosBox'))
        falso = tmpd / 'falso.xlsx'; falso.write_bytes(b'esto no es un zip' * 10)
        page.set_input_files('#impContratosFile', str(falso))
        page.wait_for_function("document.querySelector('#impContratosBox').innerText.includes('xlsx')")
        check('.xlsx corrupto: mensaje claro, sin errores de JS', 'No se pudo leer' in page.inner_text('#impContratosBox') or 'no es un .xlsx válido' in page.inner_text('#impContratosBox'), page.inner_text('#impContratosBox'))
        grande = tmpd / 'grande.csv'; grande.write_bytes(b'a' * (2 * 1024 * 1024 + 10))
        page.set_input_files('#impContratosFile', str(grande))
        page.wait_for_function("document.querySelector('#impContratosBox').innerText.includes('2 MB')")
        check('Archivo de más de 2 MB rechazado', True)
        sin_col = xlsx('sincol.xlsx', [['DEAF-050', 'X']], cab=['COD_DEAF', 'NOMBRE_CONTRATO'])
        page.set_input_files('#impContratosFile', str(sin_col))
        page.wait_for_function("document.querySelector('#impContratosBox').innerText.includes('Faltan columnas')")
        check('Faltan columnas obligatorias: lo indica', 'TIPO_UNIDAD' in page.inner_text('#impContratosBox') and 'CORREO_APROBADOR' in page.inner_text('#impContratosBox'))

        # --- subsectores: vista por contrato ---
        check('Subsectores agrupados por contrato (uno por contrato)', page.locator('#cfUnidades details.uc').count() == 4)
        c1 = '#cfUnidades details.uc[data-cid=c1]'
        check('Cada subsector es editable en su fila', page.locator(c1 + ' tr[data-id=u1] .f-nombre').count() == 1 and page.locator(c1 + ' tr[data-id=u1] .f-activo').count() == 1)
        n = len(escrituras(page))
        page.fill(c1 + ' .f-nueva', 'Aeródromo ' + XSS)
        page.click(c1 + ' .js-add-unidad')
        check('No permite un subsector repetido en el mismo contrato', 'ya tiene un subsector' in toast(page) and len(escrituras(page)) == n)
        page.fill(c1 + ' .f-nueva', 'Helipuerto Sur')
        page.click(c1 + ' .js-add-unidad')
        nuevas = esperar(page, 'unidades/', n)
        check('Agregar subsector crea la unidad activa del contrato', nuevas[-1]['data'] == {'contratoId': 'c1', 'nombre': 'Helipuerto Sur', 'activo': True}, nuevas[-1])
        n = len(escrituras(page))
        page.uncheck(c1 + ' tr[data-id=u1] .f-activo')
        page.click(c1 + ' tr[data-id=u1] .js-save')
        nuevas = esperar(page, 'unidades/u1', n)
        check('Desactivar = borrado lógico (guarda activo:false; no hay borrado)', nuevas[-1]['data']['activo'] is False and not any(w['op'] == 'delete' for w in escrituras(page)), nuevas[-1])
        check('No hay botón para borrar subsectores', page.locator('#cfUnidades .js-del, #cfUnidades .js-unidad-del').count() == 0)

        # --- carga de subsectores ---
        with page.expect_download() as dl:
            page.click('.js-plantilla[data-tipo=unidades][data-formato=xlsx]')
        ru = tmpd / 'plantilla_u.xlsx'; dl.value.save_as(ru)
        wb = openpyxl.load_workbook(ru)
        check('Plantilla de subsectores: COD_DEAF, SUBSECTOR, ACTIVO', [c.value for c in wb.worksheets[0][1]] == ['COD_DEAF', 'SUBSECTOR', 'ACTIVO'], wb.sheetnames)
        malo_u = xlsx('malo_u.xlsx', [
            ['DEAF-001', 'Aeródromo Nuevo', 'SI'],
            ['DEAF-001', 'aeródromo nuevo', 'SI'],
            ['DEAF-999', 'Fantasma', 'SI'],
            ['DEAF-001', 'Otro', 'quizá'],
            ['DEAF-001', '', 'SI'],
        ], cab=['COD_DEAF', 'SUBSECTOR', 'ACTIVO'])
        page.set_input_files('#impUnidadesFile', str(malo_u))
        page.wait_for_selector('#impUnidadesBox table')
        box = page.inner_text('#impUnidadesBox')
        check('Subsectores: errores por fila (repetido, contrato inexistente, ACTIVO, obligatorio)',
              'Subsector repetido en el archivo' in box and 'El COD_DEAF no existe' in box and 'ACTIVO debe ser SI o NO' in box and 'SUBSECTOR es obligatorio' in box and page.is_disabled('#impUnidadesBox .js-imp-ok'), box)
        bueno_u = xlsx('bueno_u.xlsx', [
            ['DEAF-001', 'Aeródromo Nuevo', 'SI'],
            ['DEAF-001', 'Helipuerto Sur', 'NO'],
            ['DEAF-001', 'Aeródromo ' + XSS, 'SI'],
            ['DEAF-002', 'Estación Cumbre', ''],
            ['DEAF-002', 'SECRETO-C2-UNIDAD', 'SI'],
        ], cab=['COD_DEAF', 'SUBSECTOR', 'ACTIVO'])
        page.set_input_files('#impUnidadesFile', str(bueno_u))
        page.wait_for_selector('#impUnidadesBox table')
        box = page.inner_text('#impUnidadesBox')
        check('Subsectores: 2 nuevos, 2 por actualizar (uno se reactiva), 1 sin cambios', '2 nuevos' in box and '2 por actualizar' in box and '1 sin cambios' in box and 'errores' not in box, box[:250])
        n = len(escrituras(page))
        page.click('#impUnidadesBox .js-imp-ok')
        page.wait_for_function('([n]) => window.__T.writes.length >= n + 4', arg=[n], timeout=10000)
        nuevas = escrituras(page)[n:]
        check('Se crean los nuevos y solo se cambia "activo" en el existente (sin borrar nada)',
              sorted(w['data']['nombre'] for w in nuevas if w['op'] == 'set') == ['Aeródromo Nuevo', 'Estación Cumbre']
              and sorted(json.dumps(w['data']) for w in nuevas if w['op'] == 'update') == ['{"activo": false}', '{"activo": true}'] and not any(w['op'] == 'delete' for w in nuevas), nuevas)
        check('Ninguna carga XSS y sin errores de JavaScript', page.evaluate('window.__T.xss') == 0 and not errores, errores)
    finally:
        browser.close()
        import shutil
        shutil.rmtree(tmpd, ignore_errors=True)


def escenario_pesos(p):
    print('\nD3. Pesos de contratos y subsectores (avance ponderado)')
    seed = base()
    seed['unidades/u1']['peso'] = 70
    seed['unidades/u1b'] = {'contratoId': 'c1', 'nombre': 'Aeródromo Segundo', 'activo': True, 'peso': 30}
    seed['oficial/avance-u1b'] = {'tipo': 'avance', 'contratoId': 'c1', 'unidadId': 'u1b', 'valor': 40, 'fechaReporte': d(-5), 'aprobadoTs': 1790000000000}
    seed['unidades/u2']['peso'] = 100
    seed['oficial/avance-u2'] = {'tipo': 'avance', 'contratoId': 'c2', 'unidadId': 'u2', 'valor': 20, 'fechaReporte': d(-5), 'aprobadoTs': 1790000000000}
    seed['contratos/c1']['peso'] = 75
    seed['contratos/c2']['peso'] = 25
    browser, page, errores = abrir(p, {'uid': 'u_admin', 'name': 'Admin Central', 'level': 'admin', 'seed': seed})
    try:
        ingresar(page)
        c1 = '.ccard[data-cid=c1]'
        check('Avance del contrato ponderado por el peso de sus subsectores (70%×62 + 30%×40 = 55%, no el simple 51%)',
              page.inner_text(c1 + ' .av-num') == '55%' and 'ponderado por peso' in page.inner_text(c1), page.inner_text(c1 + ' .av-num'))
        kpi = page.inner_text('#ctKpis .kpi >> nth=0')
        check('Avance global ponderado por el peso de los contratos (75%×55 + 25%×20 = 46%, no el simple 38%)', '46%' in kpi and 'ponderado por peso' in kpi, kpi)
        tab(page, 'config')
        d1 = '#cfUnidades details.uc[data-cid=c1]'
        check('Cada subsector muestra su peso y la suma del contrato', page.input_value(d1 + ' tr[data-id=u1] .f-peso') == '70' and 'Suma: 100% ✓' in page.inner_text(d1 + ' .js-peso-suma'), page.inner_text(d1 + ' .js-peso-suma'))
        n = len(escrituras(page))
        page.fill(d1 + ' tr[data-id=u1] .f-peso', '60')
        check('La suma se actualiza al escribir y dice cuánto falta', 'Suma: 90% · faltan 10%' in page.inner_text(d1 + ' .js-peso-suma'), page.inner_text(d1 + ' .js-peso-suma'))
        page.click(d1 + ' .js-guardar-pesos')
        check('Pesos que no suman 100% no se guardan', 'exactamente 100%' in toast(page) and len(escrituras(page)) == n)
        page.fill(d1 + ' tr[data-id=u1] .f-peso', 'abc')
        page.click(d1 + ' .js-guardar-pesos')
        check('Un peso no numérico no se guarda', 'Indica el peso' in toast(page) and len(escrituras(page)) == n)
        page.fill(d1 + ' tr[data-id=u1] .f-peso', '60,5')
        page.fill(d1 + ' tr[data-id=u1b] .f-peso', '39,5')
        check('Acepta coma decimal', 'Suma: 100% ✓' in page.inner_text(d1 + ' .js-peso-suma'))
        page.click(d1 + ' .js-guardar-pesos')
        page.wait_for_function('([n]) => window.__T.writes.length >= n + 2', arg=[n], timeout=8000)
        nuevas = escrituras(page)[n:]
        check('Se guarda solo el campo peso de cada subsector', sorted((w['path'], json.dumps(w['data'])) for w in nuevas) == [('unidades/u1', '{"peso": 60.5}'), ('unidades/u1b', '{"peso": 39.5}')], nuevas)
        tab(page, 'contratos')
        check('El avance se recalcula con los nuevos pesos (60,5%×62 + 39,5%×40 = 53%)', page.inner_text(c1 + ' .av-num') == '53%', page.inner_text(c1 + ' .av-num'))
        tab(page, 'config')
        page.click(d1 + ' .js-repartir')
        check('"Repartir en partes iguales" llena 50 y 50', page.input_value(d1 + ' tr[data-id=u1] .f-peso') == '50' and page.input_value(d1 + ' tr[data-id=u1b] .f-peso') == '50' and 'Suma: 100% ✓' in page.inner_text(d1 + ' .js-peso-suma'))
        n = len(escrituras(page))
        page.fill(d1 + ' tr[data-id=u1] .f-nombre', 'Aeródromo Principal')
        page.click(d1 + ' tr[data-id=u1] .js-save')
        nuevas = esperar(page, 'unidades/u1', n)
        check('Editar el nombre de un subsector no pierde su peso guardado', nuevas[-1]['data']['nombre'] == 'Aeródromo Principal' and nuevas[-1]['data']['peso'] == 60.5, nuevas[-1])
        # contratos
        pc = '#cfPesoContratos'
        check('Bloque de peso de cada contrato con su suma', page.input_value(pc + ' tr[data-cid=c1] .f-pesoc') == '75' and 'Suma: 100% ✓' in page.inner_text(pc + ' .js-peso-suma'))
        n = len(escrituras(page))
        page.fill(pc + ' tr[data-cid=c1] .f-pesoc', '80')
        page.click(pc + ' .js-guardar-pesos-c')
        check('Pesos de contratos que suman 105% no se guardan', 'suman 105%' in toast(page) and len(escrituras(page)) == n, toast(page))
        page.click(pc + ' .js-repartir-c')
        page.click(pc + ' .js-guardar-pesos-c')
        page.wait_for_function('([n]) => window.__T.writes.length >= n + 2', arg=[n], timeout=8000)
        nuevas = escrituras(page)[n:]
        check('Pesos de contratos guardados (50/50) sin tocar otros campos', sorted((w['path'], json.dumps(w['data'])) for w in nuevas) == [('contratos/c1', '{"peso": 50}'), ('contratos/c2', '{"peso": 50}')], nuevas)
        n = len(escrituras(page))
        page.fill('#cfContratos tr[data-id=c1] .f-nombre', 'Red Aeroportuaria Austral')
        page.click('#cfContratos tr[data-id=c1] .js-save')
        nuevas = esperar(page, 'contratos/c1', n)
        check('Guardar el formulario del contrato conserva su peso', nuevas[-1]['data'].get('peso') == 50, nuevas[-1])
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_sin_perfil(p):
    print('\nE. Colaborador sin perfil asignado')
    browser, page, errores = abrir(p, {'uid': 'u_nuevo', 'name': 'Persona Nueva', 'level': 'interact', 'seed': base()})
    try:
        page.wait_for_selector('#gateFatal', state='visible', timeout=15000)
        check('Explica que no tiene perfil y no lo deja entrar', 'no tiene un perfil' in page.inner_text('#gateFatal') and page.is_hidden('#gateBtn'))
        check('Sin errores de JavaScript', not errores, errores)
    finally:
        browser.close()


def escenario_local(p):
    print('\nF. Modo local (HTML abierto como archivo, sin claude.ai)')
    tmp = AQUI / '_tmp_prueba.html'
    tmp.write_text('<script>window.__T = { xss: 0 };</script>\n' + PAGINA, encoding='utf-8')
    browser = p.chromium.launch(executable_path=CHROME, headless=True)
    page = browser.new_page(viewport={'width': 1280, 'height': 900})
    errores = []
    page.on('pageerror', lambda e: errores.append(str(e)))
    try:
        page.goto(tmp.as_uri())
        page.wait_for_selector('#gatePerfil', state='visible', timeout=15000)
        check('Sin claude.ai arranca en modo local', 'Modo local' in page.inner_text('#gatePerfilBox'))
        page.select_option('#gateRol', 'if')
        ingresar(page)
        check('Carga los datos de ejemplo (3 contratos)', page.locator('#ctCards .ccard').count() == 3)
        check('Botón para restablecer datos visible', page.is_visible('#btnRestablecer'))
        tab(page, 'avance')
        page.select_option('#avContrato', 'dA')
        page.select_option('#cgAvUnidad', 'uA2')
        page.fill('#cgAvPct', '55')
        page.click('#formAvance button[type=submit]')
        check('Solicitud enviada en modo local', 'Enviado a revisión' in toast(page))

        page.reload()
        page.wait_for_selector('#gatePerfil', state='visible', timeout=15000)
        page.select_option('#gateRol', 'if')
        ingresar(page)
        tab(page, 'solicitudes')
        check('Los datos persisten al recargar la página', '55%' in page.inner_text('#solLista'))

        page.select_option('#verComo', 'validador')
        page.wait_for_timeout(400)
        tab(page, 'bandeja')
        card = page.locator('#bjLista .sol', has_text='55%')
        card.locator('.js-aprobar').click()
        check('Aprobación en modo local', 'Aprobada' in toast(page))
        page.select_option('#verComo', 'gabinete')
        page.wait_for_timeout(500)
        av = page.inner_text('#panel-gabinete .ccard[data-cid=dA] .av-num')
        check('Gabinete ve el nuevo avance oficial ponderado (70%×68 + 30%×55 = 64%)', av == '64%' and 'ponderado por peso' in page.inner_text('#panel-gabinete .ccard[data-cid=dA]'), av)

        page.click('#btnRestablecer')
        page.click('#btnRestablecer')
        page.wait_for_timeout(500)
        av = page.inner_text('#panel-gabinete .ccard[data-cid=dA] .av-num')
        check('Restablecer vuelve a los datos de ejemplo (70%×68 + 30%×41 = 60%)', av == '60%', av)
        check('Ninguna carga XSS y sin errores de JavaScript', page.evaluate('window.__T.xss') == 0 and not errores, errores)
    finally:
        browser.close()


def escenario_entregable(p):
    print('\nG. Entregable autocontenido "Bitacora de Obra.html" abierto desde otra carpeta')
    import shutil
    import tempfile
    origen = AQUI.parent / 'frontend' / 'Bitacora de Obra.html'
    check('El entregable existe y está regenerado (contiene la versión v10: Gantt y semáforo)', origen.exists() and 'js-g-hito' in origen.read_text(encoding='utf-8') and 'sem-fig' in origen.read_text(encoding='utf-8'))
    carpeta = pathlib.Path(tempfile.mkdtemp(prefix='bitacora_otra_'))
    copia = carpeta / 'Bitacora de Obra.html'
    shutil.copy(origen, copia)
    browser = p.chromium.launch(executable_path=CHROME, headless=True)
    page = browser.new_page(viewport={'width': 1280, 'height': 900})
    errores, consola = [], []
    page.on('pageerror', lambda e: errores.append(str(e)))
    page.on('console', lambda m: consola.append(m.text) if m.type == 'error' else None)
    # Las tipografías de Google Fonts pueden no cargar sin internet: no es un fallo de la herramienta.
    page.route('**/fonts.googleapis.com/**', lambda r: r.abort())
    try:
        page.goto(copia.as_uri())
        page.wait_for_selector('#gatePerfil', state='visible', timeout=15000)
        check('El logo de la pantalla de ingreso carga (incrustado)', page.evaluate("document.querySelector('.gate-logo').complete && document.querySelector('.gate-logo').naturalWidth > 0"))
        check('Es un solo archivo: los logos van en base64, sin rutas relativas a assets/', 'src="assets/' not in copia.read_text(encoding='utf-8') and copia.read_text(encoding='utf-8').count('data:image/png;base64,') == 2)
        page.select_option('#gateRol', 'admin')
        ingresar(page)
        check('El logo del banner carga y mide 64 px de alto', page.evaluate("document.querySelector('.appbar-logo').naturalWidth > 0 && Math.round(document.querySelector('.appbar-logo').getBoundingClientRect().height) === 64"))
        for t in ('permisos', 'config', 'bandeja'):
            tab(page, t)
        page.select_option('#verComo', 'if')
        page.wait_for_timeout(300)
        tab(page, 'hitos')
        check('Gantt y navegación funcionan desde el archivo suelto', page.locator('.js-g-hito').count() > 0)
        consola = [m for m in consola if 'fonts.googleapis' not in m and 'ERR_FAILED' not in m]
        check('Sin errores de JavaScript ni de consola', not errores and not consola, errores + consola)
    finally:
        browser.close()
        shutil.rmtree(carpeta, ignore_errors=True)


if __name__ == '__main__':
    with sync_playwright() as p:
        escenario_if(p)
        escenario_if_dos_contratos(p)
        escenario_validador(p)
        escenario_gabinete(p)
        escenario_admin(p)
        escenario_cargas_admin(p)
        escenario_pesos(p)
        escenario_sin_perfil(p)
        escenario_local(p)
        escenario_entregable(p)
    (AQUI / '_tmp_prueba.html').unlink(missing_ok=True)
    fallas = [r for r in resultados if not r[1]]
    print(f'\n{len(resultados) - len(fallas)}/{len(resultados)} pruebas OK')
    sys.exit(1 if fallas else 0)
